"""日本食品標準成分表（八訂）のExcel「第2章（データ）」を、アプリ用の foods.json に変換する。

使い方:
    python tools/convert_foods.py 成分表.xlsx foods.json

値の扱い（成分表の記号）:
    数値          そのまま
    (数値)        推定値。数値として使う
    Tr / (Tr)     微量。0 として扱う
    - / * / 空欄  未測定・記載なし。null（計算では足さない）
"""
import json
import re
import sys
import unicodedata

from openpyxl import load_workbook

SOURCE = "日本食品標準成分表（八訂）増補2023年（文部科学省）"
IDENTIFIER_LABEL = "成分識別子"

# アプリで使う成分（nutrients.js の code と同じ）
CODES = [
    "ENERC_KCAL", "PROT-", "FAT-", "CHOCDF-", "FIB-", "NACL_EQ",
    "ENERC", "WATER", "PROTCAA", "FATNLEA", "CHOLE", "CHOAVLM", "CHOAVL", "CHOAVLDF-",
    "POLYL", "OA", "ASH", "ALC",
    "NA", "K", "CA", "MG", "P", "FE", "ZN", "CU", "MN", "ID", "SE", "CR", "MO",
    "VITA_RAE", "RETOL", "CARTA", "CARTB", "CRYPXB", "CARTBEQ", "VITD",
    "TOCPHA", "TOCPHB", "TOCPHG", "TOCPHD", "VITK", "THIA", "RIBF", "NIA", "NE",
    "VITB6A", "VITB12", "FOL", "PANTAC", "BIOT", "VITC",
]

# よく使う食材の呼び名。terms をすべて含む食品のうち、食品名が最も短いものに対応させる
# "!" で始まる語は「含まないこと」の条件
ALIASES = [
    ("鶏もも肉（皮つき）", ["にわとり", "若どり", "もも", "皮つき", "生"]),
    ("鶏もも肉（皮なし）", ["にわとり", "若どり", "もも", "皮なし", "生"]),
    ("鶏むね肉（皮つき）", ["にわとり", "若どり", "むね", "皮つき", "生"]),
    ("鶏むね肉（皮なし）", ["にわとり", "若どり", "むね", "皮なし", "生"]),
    ("鶏ささみ", ["にわとり", "若どり", "ささみ", "生"]),
    ("鶏ひき肉", ["にわとり", "ひき肉", "生"]),
    ("豚ロース", ["ぶた", "大型種肉", "ロース", "脂身つき", "生"]),
    ("豚バラ", ["ぶた", "大型種肉", "ばら", "脂身つき", "生"]),
    ("豚もも（豚こま）", ["ぶた", "大型種肉", "もも", "脂身つき", "生"]),
    ("豚ひき肉", ["ぶた", "ひき肉", "生"]),
    ("牛ひき肉", ["うし", "ひき肉", "生"]),
    ("牛肩ロース（輸入）", ["うし", "輸入牛肉", "かたロース", "脂身つき", "生"]),
    ("ベーコン", ["ばらベーコン ばらベーコン"]),
    ("ロースハム", ["ロースハム", "!ゆで", "!焼き", "!フライ"]),
    ("ウインナー", ["ウインナーソーセージ", "!ゆで", "!焼き", "!フライ"]),
    ("卵（全卵・生）", ["鶏卵", "全卵", "生"]),
    ("ご飯（白米）", ["水稲めし", "精白米", "うるち米"]),
    ("ご飯（玄米）", ["水稲めし", "玄米"]),
    ("食パン", ["角形食パン", "!耳", "!焼き"]),
    ("うどん（ゆで）", ["うどん", "ゆで"]),
    ("そば（ゆで）", ["そば", "ゆで"]),
    ("スパゲッティ（ゆで）", ["スパゲッティ", "ゆで"]),
    ("中華めん（ゆで）", ["中華めん", "ゆで"]),
    ("オートミール", ["オートミール"]),
    ("木綿豆腐", ["木綿豆腐"]),
    ("絹ごし豆腐", ["絹ごし豆腐"]),
    ("納豆", ["糸引き納豆"]),
    ("牛乳", ["普通牛乳"]),
    ("ヨーグルト（無糖）", ["ヨーグルト", "全脂無糖"]),
    ("プロセスチーズ", ["プロセスチーズ"]),
    ("バター", ["無発酵バター", "有塩バター"]),
    ("鮭", ["しろさけ", "生"]),
    ("サバ", ["まさば", "生"]),
    ("ツナ缶（油漬）", ["缶詰", "油漬", "フレーク", "ライト"]),
    ("まぐろ赤身", ["くろまぐろ", "赤身", "生"]),
    ("えび（ブラックタイガー）", ["ブラックタイガー", "生"]),
    ("たまねぎ", ["たまねぎ", "りん茎", "生"]),
    ("にんじん", ["にんじん", "根", "皮つき", "生"]),
    ("じゃがいも", ["じゃがいも", "塊茎", "皮なし", "生"]),
    ("キャベツ", ["キャベツ", "結球葉", "生", "!めキャベツ", "!レッド", "!グリーンボール"]),
    ("レタス", ["レタス", "土耕栽培", "結球葉", "生"]),
    ("トマト", ["赤色トマト", "果実", "生"]),
    ("きゅうり", ["きゅうり", "果実", "生"]),
    ("ほうれん草", ["ほうれんそう", "葉", "通年平均", "生"]),
    ("ブロッコリー", ["ブロッコリー", "花序", "生"]),
    ("もやし", ["りょくとうもやし", "生"]),
    ("なす", ["なす", "果実", "生"]),
    ("長ねぎ", ["根深ねぎ", "葉", "軟白", "生"]),
    ("大根", ["だいこん", "根", "皮つき", "生"]),
    ("白菜", ["はくさい", "結球葉", "生"]),
    ("ピーマン", ["青ピーマン", "果実", "生"]),
    ("かぼちゃ", ["西洋かぼちゃ", "果実", "生"]),
    ("しいたけ", ["生しいたけ", "菌床栽培", "生"]),
    ("しめじ", ["ぶなしめじ", "生"]),
    ("バナナ", ["バナナ", "生"]),
    ("りんご", ["りんご", "皮なし", "生"]),
    ("みかん", ["うんしゅうみかん", "じょうのう", "普通", "生"]),
    ("サラダ油", ["調合油"]),
    ("オリーブ油", ["オリーブ油"]),
    ("ごま油", ["ごま油"]),
    ("砂糖（上白糖）", ["上白糖"]),
    ("しょうゆ（濃口）", ["こいくちしょうゆ"]),
    ("みそ（淡色辛みそ）", ["淡色辛みそ"]),
    ("みりん", ["本みりん"]),
    ("料理酒（清酒）", ["清酒", "普通酒"]),
    ("マヨネーズ", ["マヨネーズ", "全卵型"]),
    ("ケチャップ", ["トマトケチャップ"]),
    ("塩", ["食塩"]),
]


def normalize(s):
    """検索用に文字をそろえる（nutrients.js の normalizeText と同じ処理）"""
    s = unicodedata.normalize("NFKC", str(s or ""))
    s = "".join(chr(ord(c) - 0x60) if "ァ" <= c <= "ヶ" else c for c in s)
    return re.sub(r"[\s<>\[\]()「」『』・、，,.。/／\-－]", "", s.lower())


def parse_value(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 3)
    s = str(v).strip().strip("()")
    if s == "Tr":
        return 0
    try:
        return round(float(s), 3)
    except ValueError:
        return None  # "-", "*", 空欄


def main(src, dst):
    wb = load_workbook(src, read_only=True, data_only=True)

    # 食品群の番号 → 名前（シート名 "11肉類" などから）
    groups = {}
    for name in wb.sheetnames:
        m = re.match(r"^(\d+)(.+)$", name)
        if m:
            groups[m.group(1).zfill(2)] = m.group(2)

    ws = wb["表全体"]
    rows = ws.iter_rows(values_only=True)
    col = None
    note_col = None
    foods = []
    for row in rows:
        if col is None:
            # 「成分識別子」の行を見つけて、各成分の列番号を調べる
            if IDENTIFIER_LABEL in row:
                ids = [str(c).strip() if c else "" for c in row]
                col = {code: ids.index(code) for code in CODES}
                note_col = len(ids) - 1 if not ids[-1] else None
            continue
        food_id = row[1]
        if not food_id or not re.match(r"^\d{5}$", str(food_id)):
            continue
        name = re.sub(r"[\s　]+", " ", str(row[3])).strip()
        note = str(row[61] or "") if len(row) > 61 else ""
        alt = ""
        m = re.search(r"別名[：:]\s*([^\n]+)", note)
        if m:
            alt = m.group(1).strip()
        values = [parse_value(row[col[c]]) for c in CODES]
        foods.append([str(food_id), name, groups.get(str(row[0]).zfill(2), ""), alt] + values)

    # 呼び名を食品番号に対応させる
    aliases = []
    for label, terms in ALIASES:
        need = [normalize(t) for t in terms if not t.startswith("!")]
        ng = [normalize(t[1:]) for t in terms if t.startswith("!")]
        hits = [f for f in foods
                if all(t in normalize(f[1]) for t in need) and not any(t in normalize(f[1]) for t in ng)]
        if not hits:
            print(f"!! 見つからない: {label} {terms}")
            continue
        best = min(hits, key=lambda f: len(f[1]))
        aliases.append({"label": label, "id": best[0]})
        print(f"{label:<16} → {best[0]} {best[1]}")

    out = {"source": SOURCE, "keys": CODES, "foods": foods, "aliases": aliases}
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    print(f"食品 {len(foods)} 件、呼び名 {len(aliases)} 件を {dst} に書き出しました")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
